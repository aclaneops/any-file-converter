import os
import shutil
import io
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import traceback

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = "uploads"
OUTPUT_DIR = "outputs"
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

def pdf_to_word(input_path, output_path):
    from pdf2docx import Converter
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()

def pdf_to_excel(input_path, output_path):
    import pdfplumber
    import pandas as pd
    with pdfplumber.open(input_path) as pdf:
        all_data = []
        for page in pdf.pages:
            table = page.extract_table()
            if table:
                all_data.extend(table)
    if not all_data:
        raise Exception("No tables found in PDF")
    df = pd.DataFrame(all_data[1:], columns=all_data[0])
    df.to_excel(output_path, index=False)

def pdf_to_ppt(input_path, output_path):
    import fitz
    from pptx import Presentation
    
    prs = Presentation()
    doc = fitz.open(input_path)
    
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        temp_img_path = f"{input_path}_temp_page.png"
        pix.save(temp_img_path)
        
        slide_layout = prs.slide_layouts[6] # blank slide
        slide = prs.slides.add_slide(slide_layout)
        
        # Adding picture filling height
        slide.shapes.add_picture(temp_img_path, 0, 0, height=prs.slide_height)
        os.remove(temp_img_path)
        
    prs.save(output_path)

def image_to_word(input_path, output_path):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('Converted Image', 0)
    doc.add_picture(input_path, width=Inches(6.0))
    doc.add_paragraph('Catatan: Ekstraksi teks dari gambar membutuhkan aplikasi Tesseract OCR. Saat ini gambar dimasukkan ke dalam dokumen.')
    doc.save(output_path)

def compress_file(input_path, output_path, ext):
    ext = ext.lower()
    if ext == ".pdf":
        import fitz
        doc = fitz.open(input_path)
        doc.save(output_path, garbage=4, deflate=True)
    elif ext in [".png", ".jpg", ".jpeg"]:
        from PIL import Image
        img = Image.open(input_path)
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.save(output_path, quality=30, optimize=True)
    else:
        # Jika bukan pdf atau gambar, salin saja
        shutil.copy(input_path, output_path)

def word_to_ppt(input_path, output_path):
    from docx import Document
    from pptx import Presentation
    
    doc = Document(input_path)
    prs = Presentation()
    slide_layout = prs.slide_layouts[1] # Title and Content
    
    current_slide = None
    tf = None
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text: continue
        
        # Treat headings as new slide titles
        if para.style.name.startswith('Heading') or current_slide is None:
            current_slide = prs.slides.add_slide(slide_layout)
            current_slide.shapes.title.text = text
            tf = current_slide.placeholders[1].text_frame
            tf.text = "" # Clear placeholder
        else:
            p = tf.add_paragraph()
            p.text = text
            
    if current_slide is None:
        raise Exception("Could not find text to convert.")
        
    prs.save(output_path)

def ppt_to_word(input_path, output_path):
    from pptx import Presentation
    from docx import Document
    
    prs = Presentation(input_path)
    doc = Document()
    
    for i, slide in enumerate(prs.slides):
        doc.add_heading(f"Slide {i+1}", level=1)
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                doc.add_paragraph(shape.text.strip())
                
    doc.save(output_path)

def stub_conversion(input_path, output_path, msg="Conversion supported soon"):
    with open(output_path, "w") as f:
        f.write(f"This is a placeholder. {msg}")

@app.post("/api/convert")
async def convert_file(file: UploadFile = File(...), conversion_type: str = Form(...)):
    try:
        input_path = os.path.join(UPLOAD_DIR, file.filename)
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        base_name = os.path.splitext(file.filename)[0]
        output_filename = f"{base_name}_converted"
        
        if conversion_type == "pdf-to-word":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".docx")
            pdf_to_word(input_path, output_path)
        elif conversion_type == "pdf-to-excel":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".xlsx")
            pdf_to_excel(input_path, output_path)
        elif conversion_type == "pdf-to-ppt":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".pptx")
            pdf_to_ppt(input_path, output_path)
        elif conversion_type == "word-to-ppt":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".pptx")
            word_to_ppt(input_path, output_path)
        elif conversion_type == "ppt-to-word":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".docx")
            ppt_to_word(input_path, output_path)
        elif conversion_type == "image-to-word":
            output_path = os.path.join(OUTPUT_DIR, output_filename + ".docx")
            image_to_word(input_path, output_path)
        elif conversion_type == "compress":
            ext = os.path.splitext(file.filename)[1]
            output_path = os.path.join(OUTPUT_DIR, "compressed_" + file.filename)
            compress_file(input_path, output_path, ext)
        else:
            raise HTTPException(status_code=400, detail="Unsupported conversion type")
            
        return FileResponse(output_path, filename=os.path.basename(output_path))
        
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

# Mount frontend correctly
app.mount("/", StaticFiles(directory="public", html=True), name="public")
